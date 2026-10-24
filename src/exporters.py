"""Exporters: Excel workbook, Playwright/pytest script, traceability markdown."""
import ast
import re
from collections import defaultdict
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def slug(text: str, limit: int = 40) -> str:
    return re.sub(r"\W+", "_", text.lower()).strip("_")[:limit] or "test"


# ---------------------------------------------------------------- Excel
def _write_sheet(ws, headers: List[str], rows: List[List[Any]], widths: List[int]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"


def export_excel(path: str, requirements: List[Dict], tests: List[Dict], review: Dict) -> str:
    wb = Workbook()

    ws = wb.active
    ws.title = "Requirements"
    _write_sheet(
        ws,
        ["ID", "Title", "Description", "Type", "Priority"],
        [[r["id"], r["title"], r["description"], r["type"], r["priority"]] for r in requirements],
        [12, 35, 70, 16, 10],
    )

    ws = wb.create_sheet("Test Cases")
    _write_sheet(
        ws,
        ["Test ID", "Requirement", "Title", "Category", "Priority", "Preconditions", "Steps", "Expected Result"],
        [
            [
                t["id"], t["requirement_id"], t["title"], t["category"], t["priority"],
                t.get("preconditions", ""),
                "\n".join(f"{i}. {s}" for i, s in enumerate(t["steps"], 1)),
                t["expected_result"],
            ]
            for t in tests
        ],
        [10, 12, 40, 12, 10, 30, 55, 45],
    )

    by_req: Dict[str, List[Dict]] = defaultdict(list)
    for t in tests:
        by_req[t["requirement_id"]].append(t)
    ws = wb.create_sheet("Traceability")
    _write_sheet(
        ws,
        ["Requirement", "Title", "# Tests", "Positive", "Negative", "Boundary/Edge", "Test IDs"],
        [
            [
                r["id"], r["title"], len(by_req[r["id"]]),
                sum(t["category"] == "positive" for t in by_req[r["id"]]),
                sum(t["category"] == "negative" for t in by_req[r["id"]]),
                sum(t["category"] in ("boundary", "edge") for t in by_req[r["id"]]),
                ", ".join(t["id"] for t in by_req[r["id"]]),
            ]
            for r in requirements
        ],
        [12, 40, 9, 10, 10, 14, 40],
    )

    ws = wb.create_sheet("Review")
    _write_sheet(
        ws,
        ["Metric", "Value"],
        [
            ["Requirement coverage %", review.get("coverage_pct", "")],
            ["Reviewer quality score (1-10)", review.get("quality_score", "")],
            ["Revision passes used", review.get("revisions", "")],
            ["Requirements still flagged", ", ".join(review.get("gap_ids", [])) or "none"],
            ["Reviewer feedback", review.get("feedback", "")],
        ],
        [32, 90],
    )
    wb.save(path)
    return path


# ---------------------------------------------------------------- pytest / Playwright
SCRIPT_HEADER = '''"""Auto-generated Playwright + pytest suite.

Run:  pip install pytest-playwright && playwright install chromium
      pytest -v
Selectors marked TODO are guesses and must be aligned with the real application.
"""
import pytest
from playwright.sync_api import Page, expect

BASE_URL = "http://localhost:3000"  # TODO: point to the application under test
'''


def template_script(tests: List[Dict]) -> str:
    """Deterministic fallback: one skipped test per case with the steps as guidance comments."""
    out = [SCRIPT_HEADER]
    for t in tests:
        name = f"test_{slug(t['id'])}_{slug(t['title'], 35)}"
        steps = "\n".join(f"    # {i}. {s}" for i, s in enumerate(t["steps"], 1))
        title = t["title"].replace('"""', "'''")
        expected = t["expected_result"].replace('"""', "'''")
        out.append(
            f'''

def {name}(page: Page):
    """{t["id"]} [{t["requirement_id"]}] {title}"""
    page.goto(BASE_URL)
{steps}
    # Expected: {expected}
    pytest.skip("TODO: implement selectors and assertions")
'''
        )
    return "".join(out)


def clean_code(raw: str) -> str:
    """Strip markdown fences an LLM may add around code."""
    match = re.search(r"```(?:python)?\n(.*?)```", raw, re.S)
    return (match.group(1) if match else raw).strip()


def is_valid_python(code: str) -> bool:
    try:
        ast.parse(code)
        return True
    except SyntaxError:
        return False


def export_traceability_md(path: str, requirements: List[Dict], tests: List[Dict], review: Dict) -> str:
    by_req: Dict[str, List[Dict]] = defaultdict(list)
    for t in tests:
        by_req[t["requirement_id"]].append(t)
    lines = [
        "# Traceability Report",
        "",
        f"- Requirements: {len(requirements)}",
        f"- Test cases: {len(tests)}",
        f"- Coverage: {review.get('coverage_pct', 'n/a')}%",
        f"- Reviewer quality score: {review.get('quality_score', 'n/a')}/10",
        "",
        "| Requirement | Title | Tests |",
        "|---|---|---|",
    ]
    for r in requirements:
        ids = ", ".join(t["id"] for t in by_req[r["id"]]) or "NONE"
        lines.append(f"| {r['id']} | {r['title']} | {ids} |")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    return path
